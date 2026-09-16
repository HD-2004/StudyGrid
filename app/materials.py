"""Safe material ingestion for text, documents, URLs, and video/audio."""

from __future__ import annotations

import ipaddress
import os
import re
import socket
from dataclasses import dataclass
from html.parser import HTMLParser
from io import BytesIO
from pathlib import Path
from urllib.parse import urljoin, urlparse

import httpx
from docx import Document
from dotenv import load_dotenv
from openai import OpenAI, OpenAIError
from pypdf import PdfReader


TEXT_EXTENSIONS = {".txt", ".md", ".markdown", ".csv", ".json"}
VIDEO_EXTENSIONS = {".mp3", ".mp4", ".mpeg", ".mpga", ".m4a", ".wav", ".webm"}
SUPPORTED_EXTENSIONS = sorted(TEXT_EXTENSIONS | VIDEO_EXTENSIONS | {".pdf", ".docx"})


class MaterialError(RuntimeError):
    def __init__(self, message: str, status_code: int = 422) -> None:
        super().__init__(message)
        self.status_code = status_code


@dataclass(frozen=True)
class ExtractedMaterial:
    name: str
    kind: str
    text: str
    transcription_source: str | None = None


class _ReadableHtml(HTMLParser):
    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self._ignored_depth = 0
        self.parts: list[str] = []
        self.title = ""
        self._in_title = False

    def handle_starttag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        if tag in {"script", "style", "noscript", "svg"}:
            self._ignored_depth += 1
        elif tag == "title":
            self._in_title = True
        elif tag in {"p", "div", "section", "article", "li", "br", "h1", "h2", "h3"}:
            self.parts.append("\n")

    def handle_endtag(self, tag: str) -> None:
        if tag in {"script", "style", "noscript", "svg"} and self._ignored_depth:
            self._ignored_depth -= 1
        elif tag == "title":
            self._in_title = False

    def handle_data(self, data: str) -> None:
        if self._ignored_depth:
            return
        value = data.strip()
        if not value:
            return
        if self._in_title:
            self.title = f"{self.title} {value}".strip()
        self.parts.append(value)


def max_upload_bytes() -> int:
    load_dotenv()
    try:
        megabytes = int(os.getenv("MATERIAL_MAX_UPLOAD_MB", "25"))
    except ValueError as exc:
        raise RuntimeError("MATERIAL_MAX_UPLOAD_MB must be an integer.") from exc
    if not 1 <= megabytes <= 100:
        raise RuntimeError("MATERIAL_MAX_UPLOAD_MB must be between 1 and 100.")
    return megabytes * 1024 * 1024


def analysis_text_limit() -> int:
    load_dotenv()
    try:
        limit = int(os.getenv("MATERIAL_MAX_TEXT_CHARS", "60000"))
    except ValueError as exc:
        raise RuntimeError("MATERIAL_MAX_TEXT_CHARS must be an integer.") from exc
    if not 5_000 <= limit <= 200_000:
        raise RuntimeError("MATERIAL_MAX_TEXT_CHARS must be between 5000 and 200000.")
    return limit


def prepare_for_analysis(text: str) -> tuple[str, bool]:
    cleaned = _clean_text(text)
    limit = analysis_text_limit()
    if len(cleaned) <= limit:
        return cleaned, False
    # Preserve both the document's framing and later sections instead of
    # silently analyzing only its first pages.
    marker = "\n\n[...middle omitted for analysis...]\n\n"
    content_limit = limit - len(marker)
    head = int(content_limit * 0.72)
    tail = content_limit - head
    return f"{cleaned[:head]}{marker}{cleaned[-tail:]}", True


def _clean_text(text: str) -> str:
    text = text.replace("\x00", " ")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    cleaned = text.strip()
    if not cleaned:
        raise MaterialError("No readable text was found in this material.")
    return cleaned


def _decode_text(data: bytes) -> str:
    for encoding in ("utf-8-sig", "utf-16", "cp1252"):
        try:
            return _clean_text(data.decode(encoding))
        except UnicodeDecodeError:
            continue
    raise MaterialError("This text file uses an unsupported character encoding.")


def _extract_pdf(data: bytes) -> str:
    try:
        reader = PdfReader(BytesIO(data))
        if reader.is_encrypted:
            try:
                reader.decrypt("")
            except Exception as exc:
                raise MaterialError("Password-protected PDFs are not supported.") from exc
        pages = [(page.extract_text() or "").strip() for page in reader.pages[:250]]
    except MaterialError:
        raise
    except Exception as exc:
        raise MaterialError("The PDF could not be read. It may be damaged or image-only.") from exc
    return _clean_text("\n\n".join(page for page in pages if page))


def _extract_docx(data: bytes) -> str:
    try:
        document = Document(BytesIO(data))
        parts = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
        for table in document.tables:
            for row in table.rows:
                cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if cells:
                    parts.append(" | ".join(cells))
    except Exception as exc:
        raise MaterialError("The DOCX file could not be read. Try saving it again as DOCX.") from exc
    return _clean_text("\n".join(parts))


def _transcribe_media(data: bytes, filename: str) -> str:
    load_dotenv()
    api_key = os.getenv("OPENAI_API_KEY", "").strip()
    if not api_key:
        raise MaterialError(
            "Video and audio transcription requires OPENAI_API_KEY. "
            "DOCX, TXT, MD, PDF, and webpage imports still work without it.",
            status_code=503,
        )
    model = os.getenv("OPENAI_TRANSCRIPTION_MODEL", "gpt-4o-mini-transcribe").strip()
    timeout = float(os.getenv("OPENAI_TIMEOUT_SECONDS", "20"))
    upload = BytesIO(data)
    upload.name = Path(filename).name or "material.mp4"  # type: ignore[attr-defined]
    try:
        result = OpenAI(api_key=api_key, timeout=timeout, max_retries=1).audio.transcriptions.create(
            model=model,
            file=upload,
            response_format="text",
        )
    except (OpenAIError, TypeError, ValueError) as exc:
        raise MaterialError(
            "The video or audio could not be transcribed. Try a shorter file or paste a transcript.",
            status_code=502,
        ) from exc
    text = result if isinstance(result, str) else getattr(result, "text", "")
    return _clean_text(text)


def extract_upload(filename: str, data: bytes) -> ExtractedMaterial:
    safe_name = Path(filename or "material").name
    suffix = Path(safe_name).suffix.casefold()
    if not data:
        raise MaterialError("Choose a non-empty file to analyze.")
    if len(data) > max_upload_bytes():
        maximum = max_upload_bytes() // (1024 * 1024)
        raise MaterialError(f"Files must be {maximum} MB or smaller.", status_code=413)

    if suffix in TEXT_EXTENSIONS:
        kind = "markdown" if suffix in {".md", ".markdown"} else "text"
        return ExtractedMaterial(safe_name, kind, _decode_text(data))
    if suffix == ".pdf":
        return ExtractedMaterial(safe_name, "pdf", _extract_pdf(data))
    if suffix == ".docx":
        return ExtractedMaterial(safe_name, "docx", _extract_docx(data))
    if suffix == ".doc":
        raise MaterialError("Legacy .doc files are not supported yet. Save the file as .docx first.")
    if suffix in VIDEO_EXTENSIONS:
        return ExtractedMaterial(
            safe_name,
            "video",
            _transcribe_media(data, safe_name),
            transcription_source="ai",
        )
    raise MaterialError(
        "Unsupported file type. Use DOCX, TXT, MD, PDF, MP3, MP4, M4A, WAV, MPEG, or WEBM.",
        status_code=415,
    )


def _validate_public_url(url: str) -> str:
    parsed = urlparse(url)
    if parsed.scheme not in {"http", "https"} or not parsed.hostname:
        raise MaterialError("Enter a valid public http or https URL.")
    try:
        addresses = socket.getaddrinfo(parsed.hostname, parsed.port or 443, type=socket.SOCK_STREAM)
    except socket.gaierror as exc:
        raise MaterialError("The URL host could not be resolved.") from exc
    for address in addresses:
        ip = ipaddress.ip_address(address[4][0])
        if not ip.is_global:
            raise MaterialError("Private or local network URLs are not allowed.", status_code=403)
    return url


def extract_url(url: str) -> ExtractedMaterial:
    current = _validate_public_url(url)
    content = b""
    content_type = ""
    encoding = "utf-8"
    maximum = max_upload_bytes()
    try:
        with httpx.Client(timeout=12, follow_redirects=False, trust_env=False) as client:
            for _ in range(4):
                with client.stream(
                    "GET",
                    current,
                    headers={"User-Agent": "StudyGrid material importer/1.0"},
                ) as response:
                    if response.status_code in {301, 302, 303, 307, 308}:
                        location = response.headers.get("location")
                        if not location:
                            raise MaterialError("The URL redirected without a destination.")
                        current = _validate_public_url(urljoin(current, location))
                        continue
                    response.raise_for_status()
                    declared = response.headers.get("content-length")
                    if declared:
                        try:
                            if int(declared) > maximum:
                                raise MaterialError(
                                    "The URL content is too large to import.", status_code=413
                                )
                        except ValueError:
                            pass
                    chunks = bytearray()
                    for chunk in response.iter_bytes():
                        chunks.extend(chunk)
                        if len(chunks) > maximum:
                            raise MaterialError(
                                "The URL content is too large to import.", status_code=413
                            )
                    content = bytes(chunks)
                    content_type = response.headers.get("content-type", "").split(";", 1)[0].casefold()
                    encoding = response.encoding or "utf-8"
                    break
            else:
                raise MaterialError("The URL redirected too many times.")
    except MaterialError:
        raise
    except httpx.HTTPError as exc:
        raise MaterialError("The URL could not be downloaded.", status_code=502) from exc

    final_name = Path(urlparse(current).path).name or urlparse(current).hostname or "webpage"
    if content_type == "application/pdf" or final_name.casefold().endswith(".pdf"):
        return ExtractedMaterial(final_name, "pdf", _extract_pdf(content))
    if content_type.startswith("video/") or content_type.startswith("audio/"):
        return ExtractedMaterial(
            final_name,
            "video",
            _transcribe_media(content, final_name),
            transcription_source="ai",
        )
    if content_type in {"text/plain", "text/markdown"}:
        return ExtractedMaterial(final_name, "url", _decode_text(content))
    if content_type not in {"text/html", "application/xhtml+xml", ""}:
        raise MaterialError("This URL does not point to a readable webpage, PDF, or media file.")

    parser = _ReadableHtml()
    try:
        parser.feed(content.decode(encoding, errors="replace"))
    except (LookupError, UnicodeError) as exc:
        raise MaterialError("The webpage HTML could not be read.") from exc
    name = parser.title[:160] or final_name
    return ExtractedMaterial(name, "url", _clean_text(" ".join(parser.parts)))
