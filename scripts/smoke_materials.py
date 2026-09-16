"""Smoke checks for bounded document extraction and material upload routes."""

from __future__ import annotations

from io import BytesIO
from pathlib import Path
import sys

ROOT = Path(__file__).resolve().parent.parent
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from docx import Document
from fastapi.testclient import TestClient
from pypdf import PdfWriter
from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

from app.ai import MaterialAnalyzer
from app.api.routes import get_material_analyzer
from app.main import app
from app.materials import analysis_text_limit, extract_upload, prepare_for_analysis


def docx_bytes() -> bytes:
    output = BytesIO()
    document = Document()
    document.add_heading("Cell biology", level=1)
    document.add_paragraph("Mitosis and membrane transport")
    document.save(output)
    return output.getvalue()


def pdf_bytes() -> bytes:
    output = BytesIO()
    writer = PdfWriter()
    page = writer.add_blank_page(width=612, height=792)
    font = DictionaryObject(
        {
            NameObject("/Type"): NameObject("/Font"),
            NameObject("/Subtype"): NameObject("/Type1"),
            NameObject("/BaseFont"): NameObject("/Helvetica"),
        }
    )
    font_ref = writer._add_object(font)
    page[NameObject("/Resources")] = DictionaryObject(
        {NameObject("/Font"): DictionaryObject({NameObject("/F1"): font_ref})}
    )
    stream = DecodedStreamObject()
    stream.set_data(b"BT /F1 12 Tf 72 720 Td (Eigenvalues and eigenvectors) Tj ET")
    page[NameObject("/Contents")] = writer._add_object(stream)
    writer.write(output)
    return output.getvalue()


assert "vectors" in extract_upload("notes.txt", b"Vectors and vector spaces").text.casefold()
assert extract_upload("outline.md", b"# Thermodynamics\n- Entropy").kind == "markdown"
assert "Mitosis" in extract_upload("syllabus.docx", docx_bytes()).text
assert "Eigenvalues" in extract_upload("lecture.pdf", pdf_bytes()).text
print("[ok] TXT, MD, DOCX, and PDF extraction")

analysis_limit = analysis_text_limit()
bounded_text, was_truncated = prepare_for_analysis("A" * analysis_limit * 2)
assert was_truncated and len(bounded_text) == analysis_limit
assert "middle omitted" in bounded_text
print("[ok] long extracted text stays inside the configured analysis boundary")

app.dependency_overrides[get_material_analyzer] = lambda: MaterialAnalyzer()
try:
    with TestClient(app) as client:
        response = client.post(
            "/api/materials/upload",
            data={"subject": "Linear Algebra"},
            files={"file": ("outline.md", b"# Vectors\n# Matrices", "text/markdown")},
        )
        assert response.status_code == 200, response.text
        payload = response.json()
        assert payload["material_type"] == "markdown"
        assert payload["source"] == "fallback"
        assert payload["topics"]
        print("[ok] multipart material upload and analysis")

        legacy = client.post(
            "/api/materials/upload",
            data={"subject": "History"},
            files={"file": ("legacy.doc", b"not-a-doc", "application/msword")},
        )
        assert legacy.status_code == 422
        assert "docx" in legacy.json()["detail"].casefold()
        print("[ok] legacy DOC receives an actionable error")

        private_url = client.post(
            "/api/materials/url",
            json={"subject": "Security", "url": "http://127.0.0.1/private"},
        )
        assert private_url.status_code == 403
        print("[ok] URL importer blocks private-network SSRF")
finally:
    app.dependency_overrides.clear()

print("material smoke passed")
